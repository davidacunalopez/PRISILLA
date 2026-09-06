from __future__ import annotations

import tempfile
import unittest
import os
from datetime import date
from pathlib import Path
from unittest.mock import MagicMock, patch

from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app import config as app_config
from app.db import excel_repo, plantilla_base
from app.main import app
from app.routers import viajes as viajes_router
from app.services.exportar import viajes_csv
from app.services.fechas import semaforo
from app.services.formularios import validar
from app.services.ia import construir_contexto, preguntar_gemini
from app.services.mensajes import generar_mensaje_contadora


class LibroTemporalTest(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        raiz = Path(self.tmp.name)
        self.excel = raiz / "data" / "basedatos.xlsx"
        self.backups = raiz / "data" / "backups"
        self.exportaciones = raiz / "data" / "exportaciones"
        self.env = raiz / ".env"
        self.config = raiz / "config.json"
        self.patches = [
            patch.dict(
                os.environ,
                {"GEMINI_API_KEY": "", "GEMINI_MODEL": "", "GEMINI_ENABLED": ""},
            ),
            patch.object(app_config, "ENV_PATH", self.env),
            patch.object(app_config, "CONFIG_PATH", self.config),
            patch.object(plantilla_base, "DATA_DIR", self.excel.parent),
            patch.object(plantilla_base, "EXCEL_PATH", self.excel),
            patch.object(excel_repo, "EXCEL_PATH", self.excel),
            patch.object(excel_repo, "BACKUP_DIR", self.backups),
            patch.object(viajes_router, "EXPORT_DIR", self.exportaciones),
        ]
        for active_patch in self.patches:
            active_patch.start()
        plantilla_base._verificado = False
        plantilla_base.asegurar_libro()

    def _configurar_ruta(
        self, salida: str, llegada: str, codigo: str = "51", precio: float = 50000, moneda: str = "CRC"
    ):
        app_config.save_config(
            {
                "empresas_trabajo": ["TCC"],
                "rutas_configuradas": [
                    {
                        "salida": salida,
                        "llegada": llegada,
                        "codigo": codigo,
                        "precio": precio,
                        "moneda": moneda,
                    }
                ]
            }
        )

    def tearDown(self):
        for active_patch in reversed(self.patches):
            active_patch.stop()
        plantilla_base._verificado = True
        self.tmp.cleanup()

    def _datos_base(self):
        camion = excel_repo.insertar(
            "Camiones",
            {"Placa": "ABC123", "Marca": "Kenworth", "Estado": "Activo"},
        )
        seguro = excel_repo.insertar(
            "Seguros",
            {
                "ID_Camion": camion["ID_Camion"],
                "TipoSeguro": "Auto",
                "Periodicidad": "Trimestral",
                "FechaInicio": "2026-01-01",
                "FechaFin": "2026-06-30",
                "Moneda": "CRC",
                "Estado": "Activo",
            },
        )
        return camion, seguro

    def test_renovacion_es_atomica_y_no_superpone_fechas(self):
        _camion, seguro = self._datos_base()
        respuesta = TestClient(app).post(
            f"/seguros/{seguro['ID_Seguro']}/renovar", follow_redirects=False
        )
        self.assertEqual(respuesta.status_code, 303)
        filas = excel_repo.leer("Seguros")
        self.assertEqual(len(filas), 2)
        self.assertEqual(filas[0]["Estado"], "Renovado")
        self.assertEqual(filas[1]["FechaInicio"], "2026-07-01")
        self.assertEqual(filas[1]["FechaFin"], "2026-09-30")
        self.assertEqual(filas[1]["FechaUltimoPago"], "2026-09-30")
        self.assertEqual(filas[1]["Estado"], "Activo")

    def test_seguro_usa_fecha_fin_como_fecha_limite_de_pago(self):
        camion = excel_repo.insertar(
            "Camiones", {"Placa": "POL1", "Marca": "M", "Estado": "Activo"}
        )
        respuesta = TestClient(app).post(
            "/seguros",
            data={
                "ID_Camion": camion["ID_Camion"],
                "TipoSeguro": "Auto",
                "Periodicidad": "Mensual",
                "FechaInicio": "2026-08-01",
                "FechaFin": "2026-08-31",
                "FechaUltimoPago": "2020-01-01",
                "Estado": "Activo",
            },
            follow_redirects=False,
        )
        self.assertEqual(respuesta.status_code, 303)
        seguro = excel_repo.leer("Seguros")[0]
        self.assertEqual(seguro["FechaUltimoPago"], "2026-08-31")

        formulario = TestClient(app).get("/seguros/nuevo")
        self.assertIn("Fecha límite de pago", formulario.text)
        self.assertIn('name="FechaUltimoPago"', formulario.text)
        self.assertIn("readonly", formulario.text)
        self.assertIn(">Trimestral</option>", formulario.text)
        self.assertNotIn(">Semestral</option>", formulario.text)

    def test_respaldos_consecutivos_tienen_nombres_unicos(self):
        excel_repo.insertar("Camiones", {"Placa": "A1", "Marca": "M", "Estado": "Activo"})
        excel_repo.insertar("Camiones", {"Placa": "A2", "Marca": "M", "Estado": "Activo"})
        archivos = list(self.backups.glob("basedatos_*.xlsx"))
        self.assertEqual(len(archivos), 2)
        self.assertEqual(len({p.name for p in archivos}), 2)

    def test_solo_conserva_los_10_respaldos_mas_recientes(self):
        for numero in range(12):
            excel_repo.insertar(
                "Camiones",
                {"Placa": f"R{numero}", "Marca": "M", "Estado": "Activo"},
            )
        self.assertEqual(len(list(self.backups.glob("basedatos_*.xlsx"))), 10)

    def test_validacion_rechaza_gasolina_negativa(self):
        errores = validar(
            "Gasolina",
            {
                "Fecha": "2026-08-14",
                "ID_Camion": "",
                "ID_Chofer": "",
                "TipoCombustible": "Diesel",
                "NumeroBoleta": "B-1",
                "CantidadComprada": 20,
                "MontoTotal": -100,
                "Moneda": "CRC",
            },
        )
        self.assertIn("Litros y total pagado deben ser mayores a cero.", errores)

    def test_archivar_conserva_historial_y_permita_editarlo(self):
        camion = excel_repo.insertar(
            "Camiones", {"Placa": "HIST1", "Marca": "M", "Estado": "Activo"}
        )
        empresa = excel_repo.insertar(
            "Empresas", {"NombreEmpresa": "Puerto", "Estado": "Activo"}
        )
        chofer = excel_repo.insertar(
            "Choferes", {"Nombre": "Ana", "Estado": "Activo"}
        )
        viaje = excel_repo.insertar(
            "Viajes",
            {
                "Desde": "2026-08-14",
                "Hasta": "2026-08-15",
                "Semana": "33",
                "ID_Salida": empresa["ID_Empresa"],
                "ID_Llegada": empresa["ID_Empresa"],
                "ID_Camion": camion["ID_Camion"],
                "ID_Chofer": chofer["ID_Chofer"],
                "Estado": "Pendiente",
            },
        )
        cliente = TestClient(app)
        respuesta = cliente.post(
            f"/camiones/{camion['ID_Camion']}/archivar", follow_redirects=False
        )
        self.assertEqual(respuesta.status_code, 303)
        self.assertEqual(excel_repo.leer_por_id("Camiones", camion["ID_Camion"])["Estado"], "Inactivo")
        self.assertIsNotNone(excel_repo.leer_por_id("Viajes", viaje["ID_Viaje"]))
        edicion = cliente.get(f"/viajes/{viaje['ID_Viaje']}/editar")
        self.assertIn("HIST1 (archivado)", edicion.text)

    def test_exportacion_neutraliza_formulas(self):
        contenido = viajes_csv([{"GuiaTirManifiesto": "=HYPERLINK(\"x\")"}]).decode("utf-8-sig")
        self.assertIn("'=HYPERLINK", contenido)

    def test_camion_propone_chofer_disponible(self):
        chofer = excel_repo.insertar("Choferes", {"Nombre": "Luis", "Estado": "Activo"})
        camion = excel_repo.insertar(
            "Camiones",
            {
                "Placa": "AUTO1",
                "Marca": "M",
                "Estado": "Activo",
                "ID_ChoferPredeterminado": chofer["ID_Chofer"],
            },
        )
        pagina = TestClient(app).get("/viajes/nuevo")
        self.assertEqual(pagina.status_code, 200)
        self.assertIn(
            f'value="{camion["ID_Camion"]}" data-chofer="{chofer["ID_Chofer"]}"',
            pagina.text,
        )

    def test_filtros_de_bitacora_son_los_solicitados(self):
        pagina = TestClient(app).get("/viajes")
        for campo in (
            "desde", "hasta", "salida", "llegada", "contenedor", "chasis", "guia",
            "codigo", "empresa_trabajo", "moneda", "precio",
            "categoria", "enviada",
        ):
            self.assertIn(f'name="{campo}"', pagina.text)
        self.assertNotIn('name="precio_min"', pagina.text)
        self.assertNotIn('name="precio_max"', pagina.text)

    def test_validacion_rechaza_precio_negativo(self):
        errores = validar(
            "Viajes",
            {
                "Fecha": "2026-08-20",
                "Semana": "34",
                "ID_Salida": "inexistente",
                "ID_Llegada": "inexistente",
                "ID_Camion": "inexistente",
                "ID_Chofer": "inexistente",
                "Precio": "-1",
                "Moneda": "CRC",
                "Categoria": "Viaje completo",
                "Estado": "Pendiente",
            },
        )
        self.assertIn("El precio de la ruta debe ser un número válido y no negativo.", errores)

    def test_semana_se_calcula_desde_la_fecha(self):
        camion = excel_repo.insertar("Camiones", {"Placa": "SEM1", "Marca": "M", "Estado": "Activo"})
        empresa = excel_repo.insertar("Empresas", {"NombreEmpresa": "Patio", "Estado": "Activo"})
        chofer = excel_repo.insertar("Choferes", {"Nombre": "Eva", "Estado": "Activo"})
        self._configurar_ruta(empresa["ID_Empresa"], empresa["ID_Empresa"], "PATIO")
        respuesta = TestClient(app).post(
            "/viajes",
            data={
                "Fecha": "2026-01-05",
                "Semana": "52",
                "ID_Salida": empresa["ID_Empresa"],
                "ID_Llegada": empresa["ID_Empresa"],
                "ID_Camion": camion["ID_Camion"],
                "ID_Chofer": chofer["ID_Chofer"],
                "EmpresaTrabajo": "TCC",
                "Precio": "50000",
                "Moneda": "CRC",
                "Estado": "Pendiente",
            },
            follow_redirects=False,
        )
        self.assertEqual(respuesta.status_code, 303)
        self.assertEqual(str(excel_repo.leer("Viajes")[0]["Semana"]), "2")

    def test_factura_calcula_precio_por_litro_y_alimenta_graficos(self):
        chofer = excel_repo.insertar("Choferes", {"Nombre": "Mario", "Estado": "Activo"})
        camion = excel_repo.insertar("Camiones", {"Placa": "GAS1", "Marca": "M", "Estado": "Activo"})
        cliente = TestClient(app)
        respuesta = cliente.post(
            "/gasolina",
            data={
                "Fecha": "2026-08-15",
                "TipoCombustible": "Diesel",
                "CantidadComprada": "50",
                "ID_Camion": camion["ID_Camion"],
                "NumeroBoleta": "FAC-100",
                "ID_Chofer": chofer["ID_Chofer"],
                "MontoTotal": "35000",
                "Moneda": "CRC",
            },
            follow_redirects=False,
        )
        self.assertEqual(respuesta.status_code, 303)
        factura = excel_repo.leer("Gasolina")[0]
        self.assertEqual(factura["PrecioPorLitro"], 700)
        pagina = cliente.get("/gasolina")
        self.assertIn("Gasto asignado por chofer", pagina.text)
        self.assertIn("Mario", pagina.text)
        self.assertIn("Último mes vs anterior", pagina.text)
        self.assertIn('name="desde"', pagina.text)
        self.assertIn("del gasto filtrado", pagina.text)

    def test_graficas_gasolina_comparan_meses_y_respetan_periodo(self):
        chofer = excel_repo.insertar("Choferes", {"Nombre": "Mario", "Estado": "Activo"})
        camion = excel_repo.insertar("Camiones", {"Placa": "GRAF1", "Marca": "M", "Estado": "Activo"})
        base = {
            "TipoCombustible": "Diesel",
            "CantidadComprada": 10,
            "ID_Camion": camion["ID_Camion"],
            "ID_Chofer": chofer["ID_Chofer"],
            "Moneda": "CRC",
        }
        excel_repo.insertar(
            "Gasolina", {**base, "Fecha": "2026-01-15", "NumeroBoleta": "ENE", "MontoTotal": 10000}
        )
        excel_repo.insertar(
            "Gasolina", {**base, "Fecha": "2026-02-15", "NumeroBoleta": "FEB", "MontoTotal": 15000}
        )
        cliente = TestClient(app)
        pagina = cliente.get("/gasolina")
        self.assertIn("50.0% más", pagina.text)
        filtrada = cliente.get("/gasolina?desde=2026-02-01&hasta=2026-02-28")
        self.assertIn("FEB", filtrada.text)
        self.assertNotIn(">ENE<", filtrada.text)
        self.assertIn('value="2026-02-01"', filtrada.text)

    def test_menu_prioriza_vistas_fuertes(self):
        pagina = TestClient(app).get("/").text
        menu = pagina.split("<nav", 1)[1].split("</nav>", 1)[0]
        posiciones = [menu.index(nombre) for nombre in ("Tablero", "Viajes", "Gasolina", "Camiones", "Seguros", "Estaciones", "Choferes")]
        self.assertEqual(posiciones, sorted(posiciones))
        self.assertGreaterEqual(pagina.count("font-semibold text-accent"), 2)

    def test_marcas_se_administran_desde_configuracion(self):
        cliente = TestClient(app)
        camion = cliente.get("/camiones/nuevo")
        self.assertIn('<select name="Marca"', camion.text)
        self.assertIn("Administrar marcas", camion.text)
        self.assertIn("Agregar marca nueva", camion.text)
        configuracion = cliente.get("/configuracion")
        self.assertIn('name="marcas_camiones"', configuracion.text)

    def test_marca_nueva_se_registra_al_vuelo_desde_camiones(self):
        respuesta = TestClient(app).post(
            "/camiones",
            data={"Placa": "NEW123", "Marca": "Scania", "Estado": "Activo"},
            follow_redirects=False,
        )
        self.assertEqual(respuesta.status_code, 303)
        self.assertIn("Scania", app_config.load_config()["marcas_camiones"])
        camiones = excel_repo.leer("Camiones")
        self.assertEqual(camiones[0]["Marca"], "Scania")

    def test_error_bloqueado_muestra_aviso_visible(self):
        respuesta = TestClient(app).get("/camiones?error=bloqueado")
        self.assertIn("está abierto en otro programa", respuesta.text)

    def test_exportar_guarda_el_archivo_y_muestra_confirmacion(self):
        respuesta = TestClient(app).get("/viajes/exportar/xlsx", follow_redirects=False)
        self.assertEqual(respuesta.status_code, 303)
        archivos = list(self.exportaciones.glob("bitacora_viajes_*.xlsx"))
        self.assertEqual(len(archivos), 1)
        pagina = TestClient(app).get(respuesta.headers["location"])
        self.assertEqual(pagina.status_code, 200)
        self.assertIn(archivos[0].name, pagina.text)
        self.assertIn("data/exportaciones", pagina.text)

    def test_importar_respaldo_reemplaza_datos_y_respalda_el_libro_actual(self):
        excel_repo.insertar("Camiones", {"Placa": "ANTES", "Marca": "M", "Estado": "Activo"})
        respaldo = self.excel.read_bytes()
        excel_repo.insertar("Camiones", {"Placa": "DESPUES", "Marca": "M", "Estado": "Activo"})
        cantidad_previa = len(list(self.backups.glob("basedatos_*.xlsx")))

        respuesta = TestClient(app).post(
            "/configuracion/importar-respaldo",
            files={
                "respaldo": (
                    "respaldo.xlsx",
                    respaldo,
                    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                )
            },
            follow_redirects=False,
        )

        self.assertEqual(respuesta.status_code, 303)
        self.assertEqual([fila["Placa"] for fila in excel_repo.leer("Camiones")], ["ANTES"])
        self.assertGreater(len(list(self.backups.glob("basedatos_*.xlsx"))), cantidad_previa)

    def test_importar_archivo_invalido_no_reemplaza_los_datos(self):
        excel_repo.insertar("Camiones", {"Placa": "SEGURO", "Marca": "M", "Estado": "Activo"})
        respuesta = TestClient(app).post(
            "/configuracion/importar-respaldo",
            files={"respaldo": ("falso.xlsx", b"esto no es un xlsx", "application/octet-stream")},
        )
        self.assertEqual(respuesta.status_code, 422)
        self.assertEqual(excel_repo.leer("Camiones")[0]["Placa"], "SEGURO")

    def test_gemini_se_configura_desde_la_interfaz_sin_mostrar_la_clave(self):
        respuesta = TestClient(app).post(
            "/configuracion/ia",
            data={
                "gemini_api_key": "clave-secreta-de-prueba",
                "gemini_model": "gemini-3.6-flash",
                "gemini_enabled": "1",
            },
            follow_redirects=False,
        )
        self.assertEqual(respuesta.status_code, 303)
        self.assertIn("GEMINI_API_KEY=clave-secreta-de-prueba", self.env.read_text(encoding="utf-8"))
        pagina = TestClient(app).get("/configuracion")
        self.assertIn("Clave configurada", pagina.text)
        self.assertNotIn("clave-secreta-de-prueba", pagina.text)

    def test_preguntale_a_la_ia_es_solo_lectura_y_muestra_respuesta(self):
        app_config.save_gemini_settings(
            api_key="clave-prueba", model="gemini-3.6-flash", enabled=True
        )
        with patch("app.routers.ia.preguntar_gemini", return_value="Resumen operativo de prueba"):
            respuesta = TestClient(app).post(
                "/ia/preguntar", data={"pregunta": "Dame un resumen"}
            )
        self.assertEqual(respuesta.status_code, 200)
        self.assertIn("Resumen operativo de prueba", respuesta.text)
        self.assertIn("La IA no modifica información", respuesta.text)

    def test_preguntale_a_la_ia_exige_configuracion(self):
        respuesta = TestClient(app).post(
            "/ia/preguntar", data={"pregunta": "Dame un resumen"}
        )
        self.assertEqual(respuesta.status_code, 422)
        self.assertIn("Configure y habilite Gemini", respuesta.text)

    def test_contexto_ia_excluye_datos_personales(self):
        excel_repo.insertar(
            "Choferes",
            {
                "Nombre": "Carla",
                "Telefono": "8888-0000",
                "Licencia": "LICENCIA-SECRETA",
                "Estado": "Activo",
            },
        )
        contexto = str(construir_contexto())
        self.assertNotIn("8888-0000", contexto)
        self.assertNotIn("LICENCIA-SECRETA", contexto)
        self.assertIn("teléfonos", contexto)

    def test_cliente_gemini_recibe_el_resumen_calculado(self):
        app_config.save_gemini_settings(
            api_key="clave-prueba-valida", model="gemini-3.6-flash", enabled=True
        )
        cliente = MagicMock()
        cliente.models.generate_content.return_value.text = "Todo en orden"
        with patch("google.genai.Client", return_value=cliente):
            respuesta = preguntar_gemini("Dame un resumen")
        self.assertEqual(respuesta, "Todo en orden")
        argumentos = cliente.models.generate_content.call_args.kwargs
        self.assertIn("Resumen calculado por BITACORA", argumentos["contents"])
        self.assertEqual(argumentos["model"], "gemini-3.6-flash")
        cliente.close.assert_called_once()

    def test_configuracion_puede_probar_conexion_gemini(self):
        app_config.save_gemini_settings(
            api_key="clave-prueba-valida", model="gemini-3.6-flash", enabled=True
        )
        with patch("app.routers.configuracion.probar_gemini", return_value="gemini-3.6-flash"):
            respuesta = TestClient(app).post(
                "/configuracion/ia/probar", follow_redirects=False
            )
        self.assertEqual(respuesta.status_code, 303)
        pagina = TestClient(app).get(respuesta.headers["location"])
        self.assertIn("Conexión con Gemini verificada correctamente", pagina.text)

    def test_ruta_configurada_completa_codigo_semana_precio_y_empresa(self):
        salida = excel_repo.insertar("Empresas", {"NombreEmpresa": "San José", "Estado": "Activo"})
        llegada = excel_repo.insertar("Empresas", {"NombreEmpresa": "Limón", "Estado": "Activo"})
        camion = excel_repo.insertar("Camiones", {"Placa": "RUTA1", "Marca": "M", "Estado": "Activo"})
        chofer = excel_repo.insertar("Choferes", {"Nombre": "Sofía", "Estado": "Activo"})
        cliente = TestClient(app)
        configurada = cliente.post(
            "/configuracion/rutas",
            data={
                "ID_Salida": salida["ID_Empresa"],
                "ID_Llegada": llegada["ID_Empresa"],
                "Codigo": "51",
                "Precio": "125.50",
                "Moneda": "USD",
            },
            follow_redirects=False,
        )
        self.assertEqual(configurada.status_code, 303)
        app_config.save_config({"empresas_trabajo": ["Cliente Uno"]})

        formulario = cliente.get("/viajes/nuevo")
        self.assertIn('name="Fecha"', formulario.text)
        self.assertNotIn('name="Desde"', formulario.text)
        self.assertNotIn('name="Hasta"', formulario.text)
        self.assertIn("Cliente Uno", formulario.text)
        self.assertIn("51", formulario.text)

        respuesta = cliente.post(
            "/viajes",
            data={
                "Fecha": "2026-01-05",
                "Semana": "99",
                "ID_Salida": salida["ID_Empresa"],
                "ID_Llegada": llegada["ID_Empresa"],
                "CodigoRuta": "CODIGO-MANIPULADO",
                "EmpresaTrabajo": "Cliente Uno",
                "ID_Camion": camion["ID_Camion"],
                "ID_Chofer": chofer["ID_Chofer"],
                "Precio": "999999",
                "Moneda": "CRC",
                "Categoria": "Desvío",
                "Estado": "Pendiente",
            },
            follow_redirects=False,
        )
        self.assertEqual(respuesta.status_code, 303)
        viaje = excel_repo.leer("Viajes")[0]
        self.assertEqual(viaje["Fecha"], "2026-01-05")
        self.assertEqual(viaje["Desde"], "2026-01-05")
        self.assertEqual(str(viaje["Semana"]), "2")
        self.assertEqual(viaje["CodigoRuta"], "51")
        self.assertEqual(viaje["EmpresaTrabajo"], "Cliente Uno")
        self.assertEqual(viaje["Precio"], 125.5)
        self.assertEqual(viaje["Moneda"], "USD")
        self.assertEqual(viaje["Categoria"], "Desvío")

        filtrada = cliente.get("/viajes?moneda=USD&precio=125.50&codigo=51")
        self.assertIn("RUTA1", filtrada.text)
        precio_distinto = cliente.get("/viajes?precio=125.51")
        self.assertNotIn("RUTA1", precio_distinto.text)
        fuera = cliente.get("/viajes?moneda=CRC")
        self.assertNotIn("RUTA1", fuera.text)

    def test_viaje_sin_codigo_de_ruta_no_se_guarda(self):
        salida = excel_repo.insertar("Empresas", {"NombreEmpresa": "A", "Estado": "Activo"})
        llegada = excel_repo.insertar("Empresas", {"NombreEmpresa": "B", "Estado": "Activo"})
        camion = excel_repo.insertar("Camiones", {"Placa": "SINRUTA", "Marca": "M", "Estado": "Activo"})
        chofer = excel_repo.insertar("Choferes", {"Nombre": "Leo", "Estado": "Activo"})
        respuesta = TestClient(app).post(
            "/viajes",
            data={
                "Fecha": "2026-08-15",
                "ID_Salida": salida["ID_Empresa"],
                "ID_Llegada": llegada["ID_Empresa"],
                "ID_Camion": camion["ID_Camion"],
                "ID_Chofer": chofer["ID_Chofer"],
                "Precio": "10000",
                "Moneda": "CRC",
                "Estado": "Pendiente",
            },
        )
        self.assertEqual(respuesta.status_code, 422)
        self.assertIn("no tiene un código configurado", respuesta.text)
        self.assertEqual(excel_repo.leer("Viajes"), [])

    def test_empresa_de_trabajo_es_obligatoria_al_registrar_viaje(self):
        salida = excel_repo.insertar("Empresas", {"NombreEmpresa": "A", "Estado": "Activo"})
        llegada = excel_repo.insertar("Empresas", {"NombreEmpresa": "B", "Estado": "Activo"})
        camion = excel_repo.insertar("Camiones", {"Placa": "EMP1", "Marca": "M", "Estado": "Activo"})
        chofer = excel_repo.insertar("Choferes", {"Nombre": "Ana", "Estado": "Activo"})
        self._configurar_ruta(salida["ID_Empresa"], llegada["ID_Empresa"])
        respuesta = TestClient(app).post(
            "/viajes",
            data={
                "Fecha": "2026-08-15",
                "ID_Salida": salida["ID_Empresa"],
                "ID_Llegada": llegada["ID_Empresa"],
                "ID_Camion": camion["ID_Camion"],
                "ID_Chofer": chofer["ID_Chofer"],
                "Categoria": "Viaje completo",
                "Estado": "Pendiente",
            },
        )
        self.assertEqual(respuesta.status_code, 422)
        self.assertIn("Seleccione la empresa para la que se realiza el viaje", respuesta.text)
        self.assertEqual(excel_repo.leer("Viajes"), [])

    def test_empresa_trabajo_nueva_se_registra_al_vuelo_desde_viajes(self):
        salida = excel_repo.insertar("Empresas", {"NombreEmpresa": "A", "Estado": "Activo"})
        llegada = excel_repo.insertar("Empresas", {"NombreEmpresa": "B", "Estado": "Activo"})
        camion = excel_repo.insertar("Camiones", {"Placa": "EMP2", "Marca": "M", "Estado": "Activo"})
        chofer = excel_repo.insertar("Choferes", {"Nombre": "Ana", "Estado": "Activo"})
        self._configurar_ruta(salida["ID_Empresa"], llegada["ID_Empresa"])
        respuesta = TestClient(app).post(
            "/viajes",
            data={
                "Fecha": "2026-08-15",
                "ID_Salida": salida["ID_Empresa"],
                "ID_Llegada": llegada["ID_Empresa"],
                "ID_Camion": camion["ID_Camion"],
                "ID_Chofer": chofer["ID_Chofer"],
                "EmpresaTrabajo": "Cliente Nuevo",
                "Categoria": "Viaje completo",
                "Estado": "Pendiente",
            },
            follow_redirects=False,
        )
        self.assertEqual(respuesta.status_code, 303)
        self.assertIn("Cliente Nuevo", app_config.load_config()["empresas_trabajo"])

    def test_tarifa_actualiza_pendientes_y_congela_viajes_enviados(self):
        salida = excel_repo.insertar("Empresas", {"NombreEmpresa": "Origen", "Estado": "Activo"})
        llegada = excel_repo.insertar("Empresas", {"NombreEmpresa": "Destino", "Estado": "Activo"})
        camion = excel_repo.insertar("Camiones", {"Placa": "TAR1", "Marca": "M", "Estado": "Activo"})
        chofer = excel_repo.insertar("Choferes", {"Nombre": "Julia", "Estado": "Activo"})
        cliente = TestClient(app)
        app_config.save_config({"empresas_trabajo": ["TCC"]})
        ruta = {
            "ID_Salida": salida["ID_Empresa"],
            "ID_Llegada": llegada["ID_Empresa"],
            "Codigo": "T-01",
            "Precio": "50000",
            "Moneda": "CRC",
        }
        self.assertEqual(
            cliente.post("/configuracion/rutas", data=ruta, follow_redirects=False).status_code,
            303,
        )

        datos_viaje = {
            "Fecha": "2026-08-15",
            "ID_Salida": salida["ID_Empresa"],
            "ID_Llegada": llegada["ID_Empresa"],
            "ID_Camion": camion["ID_Camion"],
            "ID_Chofer": chofer["ID_Chofer"],
            "EmpresaTrabajo": "TCC",
            "Precio": "1",
            "Moneda": "USD",
            "Categoria": "Viaje completo",
            "Estado": "Pendiente",
        }
        self.assertEqual(
            cliente.post("/viajes", data=datos_viaje, follow_redirects=False).status_code,
            303,
        )
        primer_viaje = excel_repo.leer("Viajes")[0]
        self.assertEqual(primer_viaje["Precio"], 50000)
        self.assertEqual(primer_viaje["Moneda"], "CRC")

        ruta.update({"Precio": "725.50", "Moneda": "USD"})
        self.assertEqual(
            cliente.post("/configuracion/rutas", data=ruta, follow_redirects=False).status_code,
            303,
        )
        self.assertEqual(excel_repo.leer("Viajes")[0]["Precio"], 725.5)
        self.assertEqual(excel_repo.leer("Viajes")[0]["Moneda"], "USD")

        self.assertEqual(
            cliente.post(
                f"/viajes/{primer_viaje['ID_Viaje']}/enviar", follow_redirects=False
            ).status_code,
            303,
        )
        ruta.update({"Precio": "800", "Moneda": "USD"})
        self.assertEqual(
            cliente.post("/configuracion/rutas", data=ruta, follow_redirects=False).status_code,
            303,
        )

        datos_viaje["Categoria"] = "Desvío"
        self.assertEqual(
            cliente.post(
                f"/viajes/{primer_viaje['ID_Viaje']}", data=datos_viaje, follow_redirects=False
            ).status_code,
            303,
        )
        viaje_editado = excel_repo.leer_por_id("Viajes", primer_viaje["ID_Viaje"])
        self.assertEqual(viaje_editado["Precio"], 725.5)
        self.assertEqual(viaje_editado["Moneda"], "USD")
        self.assertTrue(viaje_editado["Enviada"])
        self.assertEqual(viaje_editado["Categoria"], "Desvío")

        datos_viaje["Fecha"] = "2026-08-16"
        datos_viaje["Categoria"] = "Viaje completo"
        self.assertEqual(
            cliente.post("/viajes", data=datos_viaje, follow_redirects=False).status_code,
            303,
        )
        segundo_viaje = excel_repo.leer("Viajes")[1]
        self.assertEqual(segundo_viaje["Precio"], 800)
        self.assertEqual(segundo_viaje["Moneda"], "USD")

    def test_ruta_heredada_sin_precio_exige_configurar_tarifa(self):
        salida = excel_repo.insertar("Empresas", {"NombreEmpresa": "A", "Estado": "Activo"})
        llegada = excel_repo.insertar("Empresas", {"NombreEmpresa": "B", "Estado": "Activo"})
        camion = excel_repo.insertar("Camiones", {"Placa": "LEG1", "Marca": "M", "Estado": "Activo"})
        chofer = excel_repo.insertar("Choferes", {"Nombre": "Leo", "Estado": "Activo"})
        app_config.save_config(
            {
                "rutas_configuradas": [
                    {"salida": salida["ID_Empresa"], "llegada": llegada["ID_Empresa"], "codigo": "LEG"}
                ]
            }
        )
        respuesta = TestClient(app).post(
            "/viajes",
            data={
                "Fecha": "2026-08-15",
                "ID_Salida": salida["ID_Empresa"],
                "ID_Llegada": llegada["ID_Empresa"],
                "ID_Camion": camion["ID_Camion"],
                "ID_Chofer": chofer["ID_Chofer"],
                "Precio": "999",
                "Moneda": "CRC",
                "Categoria": "Viaje completo",
                "Estado": "Pendiente",
            },
        )
        self.assertEqual(respuesta.status_code, 422)
        self.assertIn("no tiene un precio vigente configurado", respuesta.text)
        self.assertEqual(excel_repo.leer("Viajes"), [])

    def test_empresas_operativas_se_presentan_como_estaciones(self):
        menu = TestClient(app).get("/").text
        self.assertIn("Estaciones", menu)
        self.assertNotIn('>Empresas<', menu)
        estaciones = TestClient(app).get("/empresas")
        self.assertIn("Puntos de salida y llegada", estaciones.text)
        configuracion = TestClient(app).get("/configuracion")
        self.assertIn('name="empresas_trabajo"', configuracion.text)
        self.assertIn("Códigos de rutas", configuracion.text)

    def test_boton_registrar_viaje_es_prominente(self):
        pagina = TestClient(app).get("/viajes")
        self.assertIn("Registrar nuevo viaje", pagina.text)
        self.assertIn("shadow-lg", pagina.text)

    def test_bitacora_ofrece_columnas_configurables_y_estado_enviada(self):
        excel_repo.insertar("Viajes", {"Fecha": "2026-08-15", "Estado": "Pendiente", "Enviada": False})
        pagina = TestClient(app).get("/viajes").text
        self.assertIn("Columnas", pagina)
        self.assertIn('data-column-toggle="contenedor"', pagina)
        self.assertIn("Generar mensaje", pagina)

    def test_migracion_conserva_fecha_de_viajes_anteriores(self):
        excel_repo.insertar(
            "Viajes",
            {"Desde": "2025-12-24", "Hasta": "2025-12-26", "Semana": "52", "Estado": "Pendiente"},
        )
        wb = load_workbook(self.excel)
        ws = wb["Viajes"]
        if "Viajes" in ws.tables:
            del ws.tables["Viajes"]
        headers = [cell.value for cell in ws[1]]
        nuevas = {"Fecha", "CodigoRuta", "Precio", "Moneda", "EmpresaTrabajo", "Categoria", "Enviada"}
        for indice in sorted((i for i, nombre in enumerate(headers, 1) if nombre in nuevas), reverse=True):
            ws.delete_cols(indice)
        wb.save(self.excel)
        wb.close()

        plantilla_base.migrar_libro(self.excel, crear_backup_previo=False)
        viaje = excel_repo.leer("Viajes")[0]
        self.assertEqual(viaje["Fecha"], "2025-12-24")
        self.assertEqual(viaje["Moneda"], "CRC")
        self.assertEqual(viaje["Categoria"], "Viaje completo")
        self.assertFalse(viaje["Enviada"])

    def test_migracion_actualiza_periodicidad_y_fecha_limite_de_seguros(self):
        camion = excel_repo.insertar(
            "Camiones", {"Placa": "MIG1", "Marca": "M", "Estado": "Activo"}
        )
        seguro = excel_repo.insertar(
            "Seguros",
            {
                "ID_Camion": camion["ID_Camion"],
                "TipoSeguro": "Auto",
                "Periodicidad": "Semestral",
                "FechaInicio": "2026-01-01",
                "FechaFin": "2026-06-30",
                "FechaUltimoPago": "2026-05-15",
                "Estado": "Activo",
            },
        )

        plantilla_base.migrar_libro(self.excel, crear_backup_previo=False)

        migrado = excel_repo.leer_por_id("Seguros", seguro["ID_Seguro"])
        self.assertEqual(migrado["Periodicidad"], "Trimestral")
        self.assertEqual(migrado["FechaUltimoPago"], "2026-06-30")

    def test_mensaje_contadora_agrupa_y_marca_solo_pendientes(self):
        salida = excel_repo.insertar("Empresas", {"NombreEmpresa": "Monte Verde", "Estado": "Activo"})
        llegada = excel_repo.insertar("Empresas", {"NombreEmpresa": "MB3", "Estado": "Activo"})
        ruptura = excel_repo.insertar("Empresas", {"NombreEmpresa": "Carrusel", "Estado": "Activo"})
        comunes = {
            "Fecha": "2026-08-15",
            "Semana": "33",
            "ID_Salida": salida["ID_Empresa"],
            "ID_Llegada": llegada["ID_Empresa"],
            "CodigoRuta": "58654",
            "Categoria": "Viaje completo",
            "EmpresaTrabajo": "TCC",
            "Precio": 159437.49,
            "Moneda": "CRC",
            "Estado": "Pendiente",
            "Enviada": False,
        }
        primero = excel_repo.insertar("Viajes", comunes)
        segundo = excel_repo.insertar("Viajes", comunes)
        excel_repo.insertar(
            "Viajes",
            {
                **comunes,
                "ID_Salida": ruptura["ID_Empresa"],
                "ID_Llegada": ruptura["ID_Empresa"],
                "CodigoRuta": "74745",
                "Categoria": "Ruptura",
                "Precio": 79718.75,
            },
        )
        excel_repo.insertar("Viajes", {**comunes, "Estado": "Enviado a contadora", "Enviada": True})

        respuesta = TestClient(app).post("/viajes/generar-mensaje", data={"query": ""})
        self.assertEqual(respuesta.status_code, 200)
        self.assertIn("Factura a TCC", respuesta.text)
        self.assertIn("**58654**", respuesta.text)
        self.assertIn("2 viajes completos de Monte Verde a MB3, a 159 437,49 colones cada uno.", respuesta.text)
        self.assertIn("1 ruptura en Carrusel, a 79 718,75 colones.", respuesta.text)
        self.assertNotIn("3 viajes completos", respuesta.text)
        self.assertTrue(excel_repo.leer_por_id("Viajes", primero["ID_Viaje"])["Enviada"])
        self.assertTrue(excel_repo.leer_por_id("Viajes", segundo["ID_Viaje"])["Enviada"])

    def test_generar_mensaje_con_ia_usa_borrador_calculado(self):
        salida = excel_repo.insertar("Empresas", {"NombreEmpresa": "A", "Estado": "Activo"})
        llegada = excel_repo.insertar("Empresas", {"NombreEmpresa": "B", "Estado": "Activo"})
        viaje = excel_repo.insertar(
            "Viajes",
            {
                "Fecha": "2026-08-15",
                "Semana": "33",
                "ID_Salida": salida["ID_Empresa"],
                "ID_Llegada": llegada["ID_Empresa"],
                "CodigoRuta": "51",
                "Categoria": "Desvío",
                "EmpresaTrabajo": "TCC",
                "Precio": 100,
                "Moneda": "USD",
                "Estado": "Pendiente",
                "Enviada": False,
            },
        )
        with patch(
            "app.routers.viajes.redactar_mensaje_contadora_ia",
            return_value="Factura a TCC\n\n**51**\n1 desvío de A a B, a 100,00 dólares.",
        ) as ia:
            respuesta = TestClient(app).post(
                "/viajes/generar-mensaje", data={"query": "", "usar_ia": "1"}
            )
        self.assertEqual(respuesta.status_code, 200)
        self.assertIn("1 desvío de A a B", respuesta.text)
        self.assertIn("Factura a TCC", ia.call_args.args[0])
        self.assertTrue(excel_repo.leer_por_id("Viajes", viaje["ID_Viaje"])["Enviada"])

    def test_formato_del_mensaje_es_determinista(self):
        mensaje = generar_mensaje_contadora(
            [
                {
                    "EmpresaTrabajo": "TCC",
                    "CodigoRuta": "77",
                    "Categoria": "Ruptura",
                    "salida": "Patio",
                    "llegada": "Patio",
                    "Precio": 1000,
                    "Moneda": "CRC",
                }
            ]
        )
        self.assertEqual(mensaje, "Factura a TCC\n\n**77**\n1 ruptura en Patio, a 1 000,00 colones.")

    def test_ruta_interna_es_permitida_y_se_destaca_en_configuracion(self):
        estacion = excel_repo.insertar(
            "Empresas", {"NombreEmpresa": "Patio interno", "Estado": "Activo"}
        )
        otra = excel_repo.insertar(
            "Empresas", {"NombreEmpresa": "Otra estación", "Estado": "Activo"}
        )
        respuesta = TestClient(app).post(
            "/configuracion/rutas",
            data={
                "ID_Salida": estacion["ID_Empresa"],
                "ID_Llegada": estacion["ID_Empresa"],
                "Codigo": "R-01",
                "Precio": "20000",
                "Moneda": "CRC",
            },
            follow_redirects=False,
        )
        self.assertEqual(respuesta.status_code, 303)
        pagina = TestClient(app).get("/configuracion")
        self.assertIn("Ruptura · ruta interna", pagina.text)
        self.assertIn("R-01", pagina.text)

    def test_ruta_interna_fuerza_categoria_ruptura(self):
        estacion = excel_repo.insertar("Empresas", {"NombreEmpresa": "Bodega", "Estado": "Activo"})
        camion = excel_repo.insertar("Camiones", {"Placa": "RUP1", "Marca": "M", "Estado": "Activo"})
        chofer = excel_repo.insertar("Choferes", {"Nombre": "Nora", "Estado": "Activo"})
        self._configurar_ruta(estacion["ID_Empresa"], estacion["ID_Empresa"], "77")
        respuesta = TestClient(app).post(
            "/viajes",
            data={
                "Fecha": "2026-08-15",
                "ID_Salida": estacion["ID_Empresa"],
                "ID_Llegada": estacion["ID_Empresa"],
                "ID_Camion": camion["ID_Camion"],
                "ID_Chofer": chofer["ID_Chofer"],
                "EmpresaTrabajo": "TCC",
                "Precio": "20000",
                "Moneda": "CRC",
                "Categoria": "Desvío",
                "Estado": "Pendiente",
            },
            follow_redirects=False,
        )
        self.assertEqual(respuesta.status_code, 303)
        viaje = excel_repo.leer("Viajes")[0]
        self.assertEqual(viaje["Categoria"], "Ruptura")
        self.assertEqual(viaje["CodigoRuta"], "77")

    def test_ruptura_se_rechaza_entre_estaciones_distintas(self):
        salida = excel_repo.insertar("Empresas", {"NombreEmpresa": "A", "Estado": "Activo"})
        llegada = excel_repo.insertar("Empresas", {"NombreEmpresa": "B", "Estado": "Activo"})
        camion = excel_repo.insertar("Camiones", {"Placa": "NORUP", "Marca": "M", "Estado": "Activo"})
        chofer = excel_repo.insertar("Choferes", {"Nombre": "Pablo", "Estado": "Activo"})
        self._configurar_ruta(salida["ID_Empresa"], llegada["ID_Empresa"], "12")
        respuesta = TestClient(app).post(
            "/viajes",
            data={
                "Fecha": "2026-08-15",
                "ID_Salida": salida["ID_Empresa"],
                "ID_Llegada": llegada["ID_Empresa"],
                "ID_Camion": camion["ID_Camion"],
                "ID_Chofer": chofer["ID_Chofer"],
                "Precio": "20000",
                "Moneda": "CRC",
                "Categoria": "Ruptura",
                "Estado": "Pendiente",
            },
        )
        self.assertEqual(respuesta.status_code, 422)
        self.assertIn("Ruptura solo puede utilizarse", respuesta.text)
        self.assertEqual(excel_repo.leer("Viajes"), [])


class SemaforoTest(unittest.TestCase):
    @patch("app.services.fechas.hoy", return_value=date(2026, 1, 1))
    def test_limites_del_semaforo(self, _hoy):
        cfg = {"umbral_rojo_dias": 7, "umbral_amarillo_dias": 28}
        self.assertEqual(semaforo("2026-01-08", cfg), "rojo")
        self.assertEqual(semaforo("2026-01-09", cfg), "amarillo")
        self.assertEqual(semaforo("2026-01-29", cfg), "amarillo")
        self.assertEqual(semaforo("2026-01-30", cfg), "verde")


if __name__ == "__main__":
    unittest.main()
