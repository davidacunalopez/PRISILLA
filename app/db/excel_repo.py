"""Lectura/escritura genérica del libro Excel, con backup antes de cada guardado."""

from __future__ import annotations

import os
import shutil
import threading
import uuid
from datetime import date, datetime
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from app.config import BACKUP_DIR, EXCEL_PATH, load_config
from app.db.plantilla_base import asegurar_libro
from app.db.schema import HOJAS, columnas


class ExcelBloqueadoError(Exception):
    """El archivo está abierto en Excel u otro programa."""


_WRITE_LOCK = threading.RLock()


def _normalizar(valor):
    if valor is None:
        return ""
    if isinstance(valor, datetime):
        return valor.date().isoformat()
    if isinstance(valor, date):
        return valor.isoformat()
    if isinstance(valor, float) and valor == int(valor):
        return int(valor)
    return valor


def _abrir(solo_lectura: bool = False):
    asegurar_libro()
    try:
        return load_workbook(EXCEL_PATH, read_only=solo_lectura, data_only=False)
    except PermissionError as exc:
        raise ExcelBloqueadoError(
            "No se puede abrir basedatos.xlsx. Cierre el archivo en Excel e intente de nuevo."
        ) from exc


def _guardar(wb) -> None:
    temporal = EXCEL_PATH.with_name(f".{EXCEL_PATH.stem}_{uuid.uuid4().hex}.tmp.xlsx")
    try:
        wb.save(temporal)
        wb.close()
        verificacion = load_workbook(temporal, read_only=True)
        verificacion.close()
        os.replace(temporal, EXCEL_PATH)
    except PermissionError as exc:
        raise ExcelBloqueadoError(
            "No se pudo guardar. Cierre basedatos.xlsx en Excel e intente de nuevo."
        ) from exc
    finally:
        wb.close()
        try:
            temporal.unlink(missing_ok=True)
        except OSError:
            pass


def _limpiar_backups(conservar: int) -> None:
    if not BACKUP_DIR.exists():
        return
    archivos = sorted(BACKUP_DIR.glob("basedatos_*.xlsx"), key=lambda p: p.stat().st_mtime, reverse=True)
    for viejo in archivos[conservar:]:
        try:
            viejo.unlink()
        except OSError:
            pass


def crear_backup() -> Path:
    asegurar_libro()
    BACKUP_DIR.mkdir(parents=True, exist_ok=True)
    ts = datetime.now().strftime("%Y%m%d_%H%M%S_%f")
    destino = BACKUP_DIR / f"basedatos_{ts}.xlsx"
    try:
        shutil.copy2(EXCEL_PATH, destino)
    except PermissionError as exc:
        raise ExcelBloqueadoError(
            "No se pudo crear el respaldo. Cierre basedatos.xlsx en Excel e intente de nuevo."
        ) from exc
    conservar = int(load_config().get("backups_conservar") or 10)
    _limpiar_backups(conservar)
    return destino


def leer(hoja: str) -> list[dict]:
    asegurar_libro()
    wb = _abrir(solo_lectura=True)
    try:
        if hoja not in wb.sheetnames:
            return []
        ws = wb[hoja]
        rows = list(ws.iter_rows(values_only=True))
    finally:
        wb.close()
    if not rows:
        return []
    headers = [str(h).strip() if h is not None else "" for h in rows[0]]
    resultado = []
    for raw in rows[1:]:
        if raw is None or all(v is None or str(v).strip() == "" for v in raw):
            continue
        fila = {}
        for idx, header in enumerate(headers):
            if not header:
                continue
            fila[header] = _normalizar(raw[idx] if idx < len(raw) else "")
        resultado.append(fila)
    return resultado


def leer_por_id(hoja: str, id_valor: str) -> dict | None:
    id_col = HOJAS[hoja]["id_col"]
    for fila in leer(hoja):
        if str(fila.get(id_col, "")).strip() == str(id_valor).strip():
            return fila
    return None


def siguiente_id(hoja: str) -> str:
    prefix = HOJAS[hoja]["id_prefix"]
    id_col = HOJAS[hoja]["id_col"]
    max_n = 0
    for fila in leer(hoja):
        raw = str(fila.get(id_col, "")).strip()
        if raw.startswith(f"{prefix}-"):
            try:
                max_n = max(max_n, int(raw.split("-", 1)[1]))
            except ValueError:
                continue
    return f"{prefix}-{max_n + 1:03d}"


def _tabla_de(ws, nombre: str):
    for tabla in ws.tables.values():
        if tabla.displayName == nombre:
            return tabla
    return None


def _expandir_tabla(ws, nombre: str, n_cols: int, last_row: int) -> None:
    tabla = _tabla_de(ws, nombre)
    if tabla is None:
        return
    tabla.ref = f"A1:{get_column_letter(n_cols)}{max(1, last_row)}"


def insertar(hoja: str, datos: dict) -> dict:
    with _WRITE_LOCK:
        crear_backup()
        cols = columnas(hoja)
        id_col = HOJAS[hoja]["id_col"]
        fila = {col: datos.get(col, "") for col in cols}
        if not fila.get(id_col):
            fila[id_col] = siguiente_id(hoja)

        wb = _abrir()
        try:
            ws = wb[hoja]
            headers = [cell.value for cell in ws[1]]
            valores = [fila.get(str(header), "") for header in headers]
            last = ws.max_row
            if last == 1 or all(ws.cell(last, c).value in (None, "") for c in range(1, len(headers) + 1)):
                if last == 1:
                    ws.append(valores)
                    last_row = 2
                else:
                    for c, val in enumerate(valores, start=1):
                        ws.cell(last, c, val)
                    last_row = last
            else:
                ws.append(valores)
                last_row = ws.max_row
            _expandir_tabla(ws, hoja, len(headers), last_row)
            _guardar(wb)
        except Exception:
            wb.close()
            raise
        return fila


def actualizar(hoja: str, id_valor: str, datos: dict) -> dict:
    with _WRITE_LOCK:
        crear_backup()
        cols = columnas(hoja)
        id_col = HOJAS[hoja]["id_col"]
        wb = _abrir()
        try:
            ws = wb[hoja]
            headers = [cell.value for cell in ws[1]]
            try:
                id_idx = headers.index(id_col) + 1
            except ValueError as exc:
                raise ValueError(f"La hoja {hoja} no tiene columna {id_col}") from exc

            encontrada = None
            for row in range(2, ws.max_row + 1):
                if str(ws.cell(row, id_idx).value or "").strip() == str(id_valor).strip():
                    encontrada = row
                    break
            if encontrada is None:
                raise KeyError(f"No existe {id_col}={id_valor} en {hoja}")

            for col_name, valor in datos.items():
                if col_name not in headers or col_name == id_col:
                    continue
                col_idx = headers.index(col_name) + 1
                ws.cell(encontrada, col_idx, valor)

            actualizada = {}
            for idx, header in enumerate(headers, start=1):
                if header:
                    actualizada[header] = _normalizar(ws.cell(encontrada, idx).value)
            _expandir_tabla(ws, hoja, len(headers), ws.max_row)
            _guardar(wb)
        except Exception:
            wb.close()
            raise
        return actualizada


def actualizar_varios(hoja: str, cambios_por_id: dict[str, dict]) -> int:
    """Actualiza varias filas en un único respaldo y guardado."""
    if not cambios_por_id:
        return 0
    with _WRITE_LOCK:
        crear_backup()
        id_col = HOJAS[hoja]["id_col"]
        wb = _abrir()
        try:
            ws = wb[hoja]
            headers = [cell.value for cell in ws[1]]
            id_idx = headers.index(id_col) + 1
            actualizadas = 0
            for row in range(2, ws.max_row + 1):
                id_fila = str(ws.cell(row, id_idx).value or "").strip()
                cambios = cambios_por_id.get(id_fila)
                if not cambios:
                    continue
                for campo, valor in cambios.items():
                    if campo in headers and campo != id_col:
                        ws.cell(row, headers.index(campo) + 1, valor)
                actualizadas += 1
            _expandir_tabla(ws, hoja, len(headers), ws.max_row)
            _guardar(wb)
            return actualizadas
        except Exception:
            wb.close()
            raise


def eliminar(hoja: str, id_valor: str) -> None:
    with _WRITE_LOCK:
        crear_backup()
        id_col = HOJAS[hoja]["id_col"]
        cols = columnas(hoja)
        wb = _abrir()
        try:
            ws = wb[hoja]
            headers = [cell.value for cell in ws[1]]
            id_idx = headers.index(id_col) + 1
            for row in range(2, ws.max_row + 1):
                if str(ws.cell(row, id_idx).value or "").strip() == str(id_valor).strip():
                    ws.delete_rows(row, 1)
                    break
            else:
                raise KeyError(f"No existe {id_col}={id_valor} en {hoja}")
            last_row = max(1, ws.max_row)
            _expandir_tabla(ws, hoja, len(headers), last_row)
            _guardar(wb)
        except Exception:
            wb.close()
            raise


def actualizar_e_insertar(
    hoja: str, id_valor: str, cambios: dict, nueva_fila: dict
) -> dict:
    """Actualiza una fila e inserta otra en un único guardado atómico."""
    with _WRITE_LOCK:
        crear_backup()
        id_col = HOJAS[hoja]["id_col"]
        wb = _abrir()
        try:
            ws = wb[hoja]
            headers = [cell.value for cell in ws[1]]
            id_idx = headers.index(id_col) + 1
            row_actual = next(
                (
                    row
                    for row in range(2, ws.max_row + 1)
                    if str(ws.cell(row, id_idx).value or "").strip() == str(id_valor).strip()
                ),
                None,
            )
            if row_actual is None:
                raise KeyError(f"No existe {id_col}={id_valor} en {hoja}")
            for campo, valor in cambios.items():
                if campo in headers and campo != id_col:
                    ws.cell(row_actual, headers.index(campo) + 1, valor)

            ids = []
            prefix = HOJAS[hoja]["id_prefix"]
            for row in range(2, ws.max_row + 1):
                raw = str(ws.cell(row, id_idx).value or "")
                if raw.startswith(f"{prefix}-"):
                    try:
                        ids.append(int(raw.split("-", 1)[1]))
                    except ValueError:
                        pass
            nueva = {col: nueva_fila.get(col, "") for col in columnas(hoja)}
            nueva[id_col] = nueva.get(id_col) or f"{prefix}-{max(ids, default=0) + 1:03d}"
            ws.append([nueva.get(str(header), "") for header in headers])
            _expandir_tabla(ws, hoja, len(headers), ws.max_row)
            _guardar(wb)
            return nueva
        except Exception:
            wb.close()
            raise


def importar_respaldo(contenido: bytes) -> Path:
    """Valida, migra y activa un respaldo; conserva antes el libro actual."""
    if not contenido:
        raise ValueError("El archivo de respaldo está vacío.")
    if len(contenido) > 100 * 1024 * 1024:
        raise ValueError("El respaldo supera el límite permitido de 100 MB.")
    with _WRITE_LOCK:
        asegurar_libro()
        temporal = EXCEL_PATH.with_name(f".{EXCEL_PATH.stem}_{uuid.uuid4().hex}.import.xlsx")
        temporal.write_bytes(contenido)
        try:
            from app.db import plantilla_base

            plantilla_base.migrar_libro(temporal, crear_backup_previo=False)
            seguridad = crear_backup()
            try:
                os.replace(temporal, EXCEL_PATH)
            except PermissionError as exc:
                raise ExcelBloqueadoError(
                    "No se pudo importar. Cierre basedatos.xlsx en Excel e intente de nuevo."
                ) from exc
            plantilla_base._verificado = True
            return seguridad
        finally:
            temporal.unlink(missing_ok=True)


def mapa_display(hoja: str) -> dict[str, str]:
    """ID → etiqueta visible (placa, nombre, etc.)."""
    meta = HOJAS[hoja]
    id_col = meta["id_col"]
    display = meta["display"]
    return {
        str(f.get(id_col, "")): str(f.get(display) or f.get(id_col) or "")
        for f in leer(hoja)
    }
