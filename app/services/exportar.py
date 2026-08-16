"""Exporta filas filtradas a xlsx/csv limpio (para la contadora)."""

from __future__ import annotations

import csv
import io
from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from app.services.fechas import formato_crc, mostrar_fecha


COLUMNAS_VIAJE_EXPORT = [
    ("Fecha", "Fecha"),
    ("Semana", "Semana"),
    ("CodigoRuta", "Código de ruta"),
    ("Categoria", "Categoría"),
    ("salida", "Estación de salida"),
    ("llegada", "Estación de llegada"),
    ("EmpresaTrabajo", "Empresa para la que se trabaja"),
    ("Precio", "Precio"),
    ("Moneda", "Moneda"),
    ("placa", "Camión"),
    ("chofer", "Chofer"),
    ("Contenedor", "Contenedor"),
    ("Chasis", "Chasis"),
    ("GuiaTirManifiesto", "Guía / TIR / Manifiesto"),
    ("Enviada", "Enviada"),
]


def _celdas_viaje(fila: dict) -> list:
    valores = []
    for key, _titulo in COLUMNAS_VIAJE_EXPORT:
        val = fila.get(key, "")
        if key == "Fecha":
            val = mostrar_fecha(val) or val
        if key == "Semana" and val not in (None, ""):
            val = f"Semana {val}"
        if key == "Enviada":
            val = "Sí" if val is True or str(val).strip().lower() in {"true", "1", "si", "sí"} else "No"
        if isinstance(val, str) and val.startswith(("=", "+", "-", "@")):
            val = "'" + val
        valores.append(val if val is not None else "")
    return valores


def viajes_xlsx(filas: list[dict]) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Bitacora"
    headers = [titulo for _k, titulo in COLUMNAS_VIAJE_EXPORT]
    ws.append(headers)
    fill = PatternFill("solid", fgColor="1A2F28")
    font = Font(bold=True, color="F4EFE6")
    for cell in ws[1]:
        cell.fill = fill
        cell.font = font
        cell.alignment = Alignment(horizontal="center")
    for fila in filas:
        ws.append(_celdas_viaje(fila))
    for idx, _ in enumerate(headers, start=1):
        ws.column_dimensions[get_column_letter(idx)].width = 22
    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{max(1, len(filas) + 1)}"
    ws.freeze_panes = "A2"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def viajes_csv(filas: list[dict]) -> bytes:
    buf = io.StringIO()
    writer = csv.writer(buf, delimiter=";", lineterminator="\n")
    writer.writerow([titulo for _k, titulo in COLUMNAS_VIAJE_EXPORT])
    for fila in filas:
        writer.writerow(_celdas_viaje(fila))
    return buf.getvalue().encode("utf-8-sig")


def nombre_export(extension: str) -> str:
    ts = datetime.now().strftime("%Y%m%d_%H%M%S_%f")
    return f"bitacora_viajes_{ts}.{extension}"


def gasolina_totales_txt(monto) -> str:
    return formato_crc(monto)
